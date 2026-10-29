"""寄存器文件解析：先完整校验并预览，不改目录、不写设备、不执行公式。"""

import io
import json
import re
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import ValidationError

from .models import Access, BitField, RegisterDefinition
from .registers import validate_definitions

MAX_FILE_SIZE = 8 * 1024 * 1024
MAX_REGISTERS = 2048
MAX_ROWS = 20000
COLORS = ["blue", "purple", "cyan", "green", "amber"]


class RegisterFileError(ValueError):
    pass


def text(value) -> str:
    return str(value).strip() if value is not None else ""


def number(value, *, label: str, default=None) -> int:
    if value is None or value == "":
        if default is not None:
            return default
        raise RegisterFileError(f"{label}不能为空")
    if isinstance(value, bool):
        raise RegisterFileError(f"{label}不是有效整数")
    if isinstance(value, int) or isinstance(value, float) and value.is_integer():
        result = int(value)
    else:
        raw = text(value).replace("_", "")
        if re.fullmatch(r"0[xX][0-9a-fA-F]+", raw):
            result = int(raw, 16)
        elif re.fullmatch(r"\d+", raw):
            result = int(raw)
        else:
            raise RegisterFileError(f"{label}格式错误：{raw[:50]}，请用十进制或 0x 十六进制")
    if not 0 <= result <= 0xFFFFFFFF:
        raise RegisterFileError(f"{label}超出 32 位范围")
    return result


def access(value, name: str) -> Access:
    raw = text(value).upper().replace(" ", "").replace("/", "")
    if "RESERVED" in name.upper() or "保留" in name:
        return Access.RESERVED
    mapping = {"RW": Access.RW, "RO": Access.RO, "WO": Access.WO, "WC": Access.WC, "W1C": Access.W1C,
               "W1_PULSE": Access.W1_PULSE, "WC_INT": Access.WC_INT,
               "RESERVED": Access.RESERVED, "保留": Access.RESERVED}
    if raw not in mapping:
        raise RegisterFileError(f"字段 {name} 的类型 {text(value)!r} 暂不支持；支持 RW、RO、WO、WC、W1C、W1_PULSE、WC_INT、Reserved")
    return mapping[raw]


def bit_range(value) -> tuple[int, int]:
    raw = text(value).replace("：", ":").strip("[] ")
    if re.fullmatch(r"\d+\.0", raw):
        raw = raw[:-2]
    match = re.fullmatch(r"(\d+)(?:\s*:\s*(\d+))?", raw)
    if not match:
        raise RegisterFileError(f"位域格式错误：{raw!r}")
    msb = int(match[1])
    lsb = int(match[2] or match[1])
    if not 0 <= lsb <= msb <= 31:
        raise RegisterFileError(f"位域越界或高低位颠倒：{raw}")
    return msb, lsb


def normalize_register(item: dict, warnings: list[str]) -> RegisterDefinition:
    address = number(item.get("address"), label="寄存器地址")
    fields = []
    covered = 0
    for index, raw in enumerate(item.get("fields", [])):
        if not isinstance(raw, dict):
            raise RegisterFileError("字段必须为对象")
        name = text(raw.get("name"))
        msb, lsb = bit_range(f"{raw.get('msb')}:{raw.get('lsb')}")
        field_access = access(raw.get("access"), name)
        field_reset = number(raw.get("reset"), label=f"字段 {name} 初始值") if raw.get("reset") not in (None, "") else None
        if field_reset is not None and field_reset >= 1 << (msb - lsb + 1):
            raise RegisterFileError(f"字段 {name} 初始值超出位宽")
        bit_field = BitField(name=name, msb=msb, lsb=lsb, access=field_access,
                            description=text(raw.get("description")),
                            color=raw.get("color") if raw.get("color") in COLORS + ["gray"] else (
                                "gray" if field_access in (Access.RO, Access.RESERVED) else COLORS[index % len(COLORS)]),
                            reset=field_reset)
        if covered & bit_field.mask:
            raise RegisterFileError(f"0x{address:08X} 字段 {name} 与其他位域重叠")
        covered |= bit_field.mask
        fields.append(bit_field)
    missing = (~covered) & 0xFFFFFFFF
    if missing:
        warnings.append(f"0x{address:08X} 未定义位已标为保留只读，不会默认允许写入")
        bit = 31
        while bit >= 0:
            if not missing & (1 << bit):
                bit -= 1
                continue
            msb = bit
            while bit >= 0 and missing & (1 << bit):
                bit -= 1
            fields.append(BitField(name="RESERVED", msb=msb, lsb=bit + 1, access=Access.RESERVED,
                                   description="未提供字段定义，禁止写入", color="gray"))
    return RegisterDefinition(address=f"0x{address:08X}", name=text(item.get("name")) or f"REG_{address:08X}",
                              group=text(item.get("group")) or "导入寄存器", description=text(item.get("description")),
                              reset=number(item.get("reset"), label="初始值", default=0),
                              fields=sorted(fields, key=lambda field: -field.msb),
                              source_version=item.get("source_version"))


def parse_json(content: bytes) -> list[dict]:
    try:
        data = json.loads(content.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise RegisterFileError("JSON 文件编码或格式错误，请使用 UTF-8") from exc
    if isinstance(data, dict):
        data = data.get("registers")
    if not isinstance(data, list) or not data:
        raise RegisterFileError("JSON 应为非空寄存器数组或包含 registers 数组的对象")
    result = []
    for index, item in enumerate(data, start=1):
        if not isinstance(item, dict):
            raise RegisterFileError(f"第 {index} 项不是寄存器对象")
        fields = item.get("fields", item.get("bitfields", item.get("bit_fields", [])))
        if not isinstance(fields, list) or len(fields) > 32:
            raise RegisterFileError(f"第 {index} 项位域必须是最多 32 项的数组")
        converted = []
        for field in fields:
            if not isinstance(field, dict):
                raise RegisterFileError(f"第 {index} 项字段格式错误")
            converted.append({"name": field.get("name"), "msb": field.get("msb", field.get("start_bit")),
                              "lsb": field.get("lsb", field.get("end_bit")),
                              "access": field.get("access", field.get("type")), "description": field.get("description"), "color": field.get("color"),
                              "reset": field.get("reset")})
        result.append({"address": item.get("address"), "name": item.get("name", item.get("description")),
                       "group": item.get("group"), "description": item.get("description"),
                       "reset": item.get("reset", item.get("init_value", item.get("data", item.get("value32bit", 0)))),
                       "fields": converted, "source_version": item.get("source_version")})
    return result


def open_excel(content: bytes):
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if len(archive.infolist()) > 3000 or sum(item.file_size for item in archive.infolist()) > 64 * 1024 * 1024:
                raise RegisterFileError("Excel 解压后内容过大，请拆分表格")
        return load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except RegisterFileError:
        raise
    except Exception as exc:
        raise RegisterFileError("无法打开 Excel，请提供有效的 .xlsx 文件") from exc


def sheet_header(sheet) -> tuple[dict | None, str, object]:
    base = None
    for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=10, max_col=64, values_only=True), start=1):
        if row_number == 2:
            base = row[1] if len(row) > 1 else None
        columns = {}
        resets = []
        for index, value in enumerate(row):
            key = text(value)
            if key == "初始值":
                resets.append(index)
            elif key:
                columns.setdefault(key, index)
        if "名称" not in columns or not ("偏移地址" in columns or "地址" in columns):
            continue
        member_index = columns.get("成员变量", 64)
        # 两列同名“初始值”必须按寄存器区/字段区分别映射，不能被后列覆盖。
        columns["register_reset"] = next((index for index in resets if index < member_index), None)
        columns["field_reset"] = next((index for index in resets if index > member_index), None)
        columns["header_row"] = row_number
        version = "v2" if "合并" in columns else "v1" if "成员变量" in columns else "simple"
        return columns, version, base
    return None, "unsupported", base


def inspect_excel(content: bytes) -> dict:
    if not content or len(content) > MAX_FILE_SIZE:
        raise RegisterFileError("寄存器文件为空或超过 8 MB")
    workbook = open_excel(content)
    try:
        sheets = []
        for sheet in workbook:
            _, version, _ = sheet_header(sheet)
            sheets.append({"name": sheet.title, "format": version})
        return {"sheets": sheets}
    except Exception as exc:
        raise RegisterFileError("无法识别工作表结构") from exc
    finally:
        workbook.close()


def parse_excel(content: bytes, warnings: list[str], selected_sheets: list[str] | None = None) -> list[dict]:
    workbook = open_excel(content)
    result = []
    total_rows = 0
    try:
        if selected_sheets is not None and (not selected_sheets or set(selected_sheets) - set(workbook.sheetnames)):
            raise RegisterFileError("请选择存在的工作表")
        for sheet in workbook.worksheets:
            if selected_sheets is not None and sheet.title not in selected_sheets:
                continue
            header, version, base = sheet_header(sheet)
            if header is None:
                if selected_sheets is not None:
                    raise RegisterFileError(f"工作表 {sheet.title} 不是支持的寄存器结构")
                warnings.append(f"已跳过非寄存器工作表：{sheet.title}")
                continue
            if sheet.max_row and sheet.max_row > MAX_ROWS or sheet.max_column and sheet.max_column > 64:
                raise RegisterFileError(f"工作表 {sheet.title} 过大，请限制为 {MAX_ROWS} 行和 64 列以内")
            current = None
            if "偏移地址" in header:
                base = number(base, label=f"{sheet.title} B2 基地址")
            for row_number, row in enumerate(sheet.iter_rows(min_row=header["header_row"] + 1, values_only=True), start=header["header_row"] + 1):
                total_rows += 1
                if total_rows > MAX_ROWS:
                    raise RegisterFileError(f"Excel 总行数超过 {MAX_ROWS}，请拆分")
                if not any(value is not None and value != "" for value in row):
                    continue

                def cell(column, default=None):
                    position = header.get(column)
                    value = row[position] if position is not None and position < len(row) else default
                    if isinstance(value, str) and value.startswith("="):
                        raise RegisterFileError(f"{sheet.title} 第 {row_number} 行包含公式，请先转换为值")
                    return value

                offset = cell("偏移地址") if "偏移地址" in header else cell("地址")
                if offset is not None and text(offset):
                    name = text(cell("名称"))
                    if not name:
                        raise RegisterFileError(f"{sheet.title} 第 {row_number} 行缺少寄存器名称")
                    address = number(offset, label=f"{sheet.title} 第 {row_number} 行地址") + (base if "偏移地址" in header else 0)
                    current = {"address": address, "name": name, "group": sheet.title,
                               "description": text(cell("说明", cell("描述", name))) or name,
                               "reset": cell("register_reset", 0), "fields": [], "source_version": version}
                    result.append(current)
                    if len(result) > MAX_REGISTERS:
                        raise RegisterFileError(f"一次最多导入 {MAX_REGISTERS} 个寄存器")
                elif text(cell("名称")):
                    raise RegisterFileError(f"{sheet.title} 第 {row_number} 行寄存器缺少地址")
                member = cell("成员变量")
                bits = cell("位域")
                if text(member) or text(bits):
                    if current is None or not text(member) or not text(bits):
                        raise RegisterFileError(f"{sheet.title} 第 {row_number} 行字段不完整或不属于寄存器")
                    msb, lsb = bit_range(bits)
                    current["fields"].append({"name": text(member), "msb": msb, "lsb": lsb,
                                               "access": cell("类型"), "description": cell("成员描述", ""),
                                               "reset": cell("field_reset")})
    except RegisterFileError:
        raise
    except Exception as exc:
        raise RegisterFileError("Excel 内容无法解析，请检查工作表和单元格格式") from exc
    finally:
        workbook.close()
    return result


def parse_register_file(filename: str, content: bytes, selected_sheets: list[str] | None = None) -> dict:
    if not content or len(content) > MAX_FILE_SIZE:
        raise RegisterFileError("寄存器文件为空或超过 8 MB")
    warnings = []
    extension = filename.rsplit(".", 1)[-1].lower()
    try:
        if extension == "json":
            items = parse_json(content)
        elif extension == "xlsx":
            items = parse_excel(content, warnings, selected_sheets)
        else:
            raise RegisterFileError("仅支持 .xlsx 和 .json 寄存器文件")
        if not items or len(items) > MAX_REGISTERS:
            raise RegisterFileError(f"文件中没有寄存器或超过 {MAX_REGISTERS} 项")
        registers = [normalize_register(item, warnings) for item in items]
        validate_definitions(registers)
    except RegisterFileError:
        raise
    except (ValidationError, ValueError, TypeError, RecursionError) as exc:
        raise RegisterFileError(f"寄存器定义无效：{str(exc)[:300]}") from exc
    return {"registers": [item.model_dump(mode="json") for item in registers],
            "warnings": warnings[:50], "count": len(registers),
            "formats": {item.group: item.source_version or "json" for item in registers},
            "groups": list(dict.fromkeys(item.group for item in registers))}


def excel_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "EXAMPLE"
    sheet.append(["TestStudio 寄存器示例，仅用于模拟；B2 为基地址，第 4 行为表头"])
    sheet.append(["基地址", "0x00002000"])
    sheet.append(["支持 RW、RO、WO、WC、W1C、W1_PULSE、WC_INT、Reserved；WC/W1_PULSE/WC_INT 等价，写 1 后硬件自动清 0"])
    sheet.append(["名称", "偏移地址", "初始值", "成员变量", "位域", "类型", "成员描述"])
    sheet.append(["EXAMPLE_CTRL", "0x00", "0x00000001", "RESERVED", "31:8", "Reserved", "保留，禁止修改"])
    sheet.append([None, None, None, "GAIN", "7:1", "RW", "增益控制"])
    sheet.append([None, None, None, "ENABLE", "0", "RW", "使能"])
    sheet.append(["EXAMPLE_STATUS", "0x04", "0x00000001", "STATUS", "31:0", "RO", "只读状态"])
    for cell in sheet[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="394BB7")
    for row in sheet.iter_rows(min_row=5):
        for cell in row:
            cell.font = Font(color="0000FF" if cell.column <= 6 else "000000")
    for column, width in {"A": 25, "B": 18, "C": 18, "D": 20, "E": 12, "F": 14, "G": 40}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A5"
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
