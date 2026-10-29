import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from backend.app import create_app
from backend.models import ConnectionCreate, RegisterDelete, RegisterImport, RegisterWrite
from backend.register_import import RegisterFileError, excel_template, inspect_excel, parse_register_file
from backend.repository import Repository
from backend.service import ServiceError, StudioService


def workbook_bytes(rows=None, base='0x00002000'):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'BLE_System'
    sheet.append(['寄存器示例'])
    sheet.append(['基地址', base])
    sheet.append([])
    sheet.append(['名称', '偏移地址', '初始值', '成员变量', '位域', '类型', '成员描述'])
    for row in rows or [
        ['IMPORTED_CTRL', '0x00', '0x00000001', 'ENABLE', '[0:0]', 'RW', '使能'],
        [None, None, None, 'MODE', '3:1', 'RW', '模式'],
        ['IMPORTED_STATUS', '0x04', '0x00000005', 'STATUS', '31:0', 'RO', '状态'],
    ]:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def parsed_registers():
    return parse_register_file('registers.xlsx', workbook_bytes())['registers']


def test_excel_reference_format_and_reserved_gaps():
    preview = parse_register_file('map.xlsx', workbook_bytes(base='0x0000_2000'))
    assert preview['count'] == 2
    assert preview['groups'] == ['BLE_System']
    first, second = preview['registers']
    assert first['address'] == '0x00002000'
    assert first['reset'] == 1
    assert first['fields'][0]['access'] == 'Reserved'
    assert first['fields'][0]['msb'] == 31
    assert first['fields'][0]['lsb'] == 4
    assert second['fields'][0]['access'] == 'RO'
    assert preview['warnings']


def test_excel_simple_absolute_format_and_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['名称', '地址', '初始值', '描述'])
    sheet.append(['NO_FIELDS', '0x00003000', 0, '只有地址的寄存器'])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    preview = parse_register_file('simple.xlsx', output.getvalue())
    assert preview['registers'][0]['fields'][0]['access'] == 'Reserved'
    assert preview['registers'][0]['fields'][0]['msb'] == 31
    assert parse_register_file('template.xlsx', excel_template())['count'] == 2


@pytest.mark.parametrize('rows,error', [
    ([['CTRL', '0x00', 0, 'A', '3:0', 'RW', ''], [None, None, None, 'B', '7:2', 'RW', '']], '重叠'),
    ([['CTRL', '0x00', 0, 'A', '32:0', 'RW', '']], '越界'),
    ([['CTRL', '0x00', 0, 'A', '0:7', 'RW', '']], '颠倒'),
    ([['CTRL', '0x00', 0, 'A', '7:0', 'W0C', '']], '暂不支持'),
    ([['CTRL', '0x02', 0, 'A', '31:0', 'RW', '']], '对齐'),
    ([['CTRL', '0x00', 0, 'A', '31:0', 'RW', ''], ['DUP', '0x00', 0, 'B', '31:0', 'RW', '']], '重复'),
    ([['CTRL', '0x00', '=1+1', 'A', '31:0', 'RW', '']], '公式'),
])
def test_excel_invalid_definitions_are_explicit(rows, error):
    with pytest.raises(RegisterFileError, match=error):
        parse_register_file('bad.xlsx', workbook_bytes(rows=rows))


def test_json_old_and_new_formats_unknown_fields_safe():
    original = parsed_registers()
    preview = parse_register_file('export.json', json.dumps({'version': 1, 'registers': original}).encode())
    assert preview['registers'] == original
    legacy = [{'address': '0x4000', 'data': '0x80000000', 'description': 'LEGACY'}]
    preview = parse_register_file('legacy.json', json.dumps(legacy).encode())
    assert preview['registers'][0]['reset'] == 0x80000000
    assert preview['registers'][0]['fields'][0]['access'] == 'Reserved'
    for filename, content in [('bad.json', b'no json'), ('bad.xlsx', b'not a zip'), ('bad.csv', b'a,b'), ('empty.json', b'[]')]:
        with pytest.raises(RegisterFileError):
            parse_register_file(filename, content)


def test_import_delete_persist_and_do_not_write_or_reset_device(tmp_path, monkeypatch):
    path = tmp_path / 'studio.db'
    service = StudioService(path)
    try:
        connection_id = service.snapshot()['connections'][0]['id']
        service.connect(connection_id)
        service.write_register(connection_id, RegisterWrite(address='0x1000', value=0, mask=1))
        driver = service.drivers[connection_id]
        writes = []
        monkeypatch.setattr(driver, 'write', lambda *args: writes.append(args))
        request = RegisterImport(mode='append', expected_revision=service.catalog_revision, registers=parsed_registers())
        service.import_registers(request)
        assert len(service.register_catalog()['registers']) == 10
        assert driver.read(0x1000) == 0x100
        assert service.read_register(connection_id, 0x2000)['value'] == 1
        with pytest.raises(ServiceError, match='发生变化'):
            service.import_registers(request)
        request.expected_revision = service.catalog_revision
        with pytest.raises(ServiceError, match='重复'):
            service.import_registers(request)
        service.delete_registers(RegisterDelete(expected_revision=service.catalog_revision, addresses=['0x2000']))
        assert driver.read(0x2000) == 1
        assert not writes
        with pytest.raises(ServiceError, match='未定义'):
            service.read_register(connection_id, 0x2000)
        other = StudioService(path)
        try:
            assert len(other.definitions) == 9
            assert 0x2000 not in other.definitions
            assert 0x2004 in other.definitions
        finally:
            other.close()
    finally:
        service.close()


def test_delete_all_stays_empty_after_restart_and_task_guard(tmp_path):
    path = tmp_path / 'empty.db'
    service = StudioService(path)
    try:
        connection_id = service.snapshot()['connections'][0]['id']
        service.connect(connection_id)
        service.start_scan(connection_id, 3)
        request = RegisterDelete(expected_revision=service.catalog_revision, addresses=['0x1000'])
        with pytest.raises(ServiceError, match='任务运行中'):
            service.delete_registers(request)
        with pytest.raises(ServiceError, match='任务运行中'):
            service.import_registers(RegisterImport(expected_revision=service.catalog_revision, registers=parsed_registers()))
        service.close()
        service.delete_registers(RegisterDelete(expected_revision=service.catalog_revision,
                                               addresses=[register.address for register in service.definitions.values()]))
        assert service.definitions == {}
    finally:
        service.close()
    other = StudioService(path)
    try:
        assert other.definitions == {}
        other.connect(connection_id)
        with pytest.raises(ServiceError, match='表为空'):
            other.start_scan(connection_id, 1)
    finally:
        other.close()


def test_clear_logs_removes_saved_records_and_new_logs_continue(tmp_path):
    service = StudioService(tmp_path / 'clear.db')
    try:
        old_ids = [entry['id'] for entry in service.snapshot()['logs']]
        assert service.clear_logs()['deleted'] >= 4
        assert service.snapshot()['logs'] == []
        assert service.repository.logs() == []
        connection_id = service.snapshot()['connections'][0]['id']
        service.connect(connection_id)
        logs = service.snapshot()['logs']
        assert len(logs) == 1
        assert logs[0]['id'] > max(old_ids)
        assert '模拟连接已建立' in logs[0]['message']
    finally:
        service.close()


def test_api_preview_commit_delete_clear_and_atomic_error(tmp_path):
    path = tmp_path / 'api.db'
    with TestClient(create_app(path)) as client:
        assert client.delete('/api/logs').status_code == 401
        client.post('/api/bootstrap')
        catalog = client.get('/api/registers/definitions').json()
        revision = catalog['revision']
        preview = client.post('/api/registers/preview?filename=map.xlsx', content=workbook_bytes()).json()
        assert len(client.get('/api/registers/definitions').json()['registers']) == 8
        invalid = client.post('/api/registers/preview?filename=bad.xlsx', content=b'bad data')
        assert invalid.status_code == 400
        assert client.get('/api/registers/definitions').json()['revision'] == revision
        result = client.post('/api/registers/import', json={'mode': 'append', 'expected_revision': revision, 'registers': preview['registers']})
        assert result.status_code == 200
        assert client.post('/api/registers/delete', json={'expected_revision': revision, 'addresses': ['0x2000']}).status_code == 409
        assert client.get('/api/registers/template').headers['content-type'].startswith('application/vnd.openxmlformats')
        assert client.delete('/api/logs').status_code == 200
        assert client.get('/api/state').json()['logs'] == []
    assert Repository(path).logs() == []


def versioned_workbook():
    """独立合成 v1/v2 结构样例，不复制芯片文件的定义或描述。"""
    workbook = Workbook()
    first = workbook.active
    first.title = 'V1_EXAMPLE'
    first.append(['模块名称', 'v1_sample'])
    first.append(['基地址', '0x0000_5000'])
    first.append(['地址位宽', 8, '数据位宽', 32])
    first.append(['名称', '说明', '偏移地址', '初始值', '成员变量', '位域', '类型', '初始值', '成员描述'])
    first.append(['V1_CTRL', 'v1 整字初始值不能被末尾字段初值覆盖', '0x00', '0x00001203', 'RESERVED', '31:16', 'RO', '0x0', '保留'])
    first.append([None, None, None, None, 'VALUE', '15:8', 'RW', '0x12', '普通字段'])
    first.append([None, None, None, None, 'RESERVED', '7:2', 'RO', '0x0', '保留'])
    first.append([None, None, None, None, 'ENABLE', '1:0', 'RW', '0x3', '使能'])
    first.append(['V1_EVENT', '特殊写语义', '0x04', '0x00000002', 'PULSE', '0:0', 'W1_PULSE', '0x0', '脉冲'])
    first.append([None, None, None, None, 'IRQ', '1:1', 'WC_INT', '0x1', '中断'])
    second = workbook.create_sheet('V2_EXAMPLE')
    second.append(['模块名称', 'v2_sample'])
    second.append(['基地址', '0x00006000'])
    second.append(['地址位宽', 16, '数据位宽', 32])
    second.append(['名称', '说明', '合并', '偏移地址', '初始值', '成员变量', '位域', '类型', '初始值', 'DFT值', '成员描述'])
    second.append(['V2_CTRL', 'v2 描述', 'true', '0x0000', '0x00101004', 'UPPER', '31:16', 'RW', '0x10', '0x20', '高位参数'])
    second.append([None, None, None, None, None, 'LOWER', '15:0', 'RW', '0x1004', '0x2004', '低位参数'])
    second.append(['V2_STATUS', None, 'false', '0x0004', '0x00000005', 'STATUS', '31:0', 'RO', '0x5', 'reserved', '只读状态'])
    notes = workbook.create_sheet('说明')
    notes.append(['这不是寄存器工作表'])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_v1_v2_select_sheets_keep_distinct_resets_and_ignore_merge_dft():
    content = versioned_workbook()
    assert inspect_excel(content)['sheets'] == [
        {'name': 'V1_EXAMPLE', 'format': 'v1'}, {'name': 'V2_EXAMPLE', 'format': 'v2'},
        {'name': '说明', 'format': 'unsupported'},
    ]
    preview = parse_register_file('mixed.xlsx', content, ['V1_EXAMPLE', 'V2_EXAMPLE'])
    assert preview['count'] == 4
    assert preview['formats'] == {'V1_EXAMPLE': 'v1', 'V2_EXAMPLE': 'v2'}
    v1, event, v2, status = preview['registers']
    assert v1['reset'] == 0x1203  # 不能误用同名的字段初始值列
    assert v1['description'] == 'v1 整字初始值不能被末尾字段初值覆盖'
    assert next(field for field in v1['fields'] if field['name'] == 'VALUE')['reset'] == 0x12
    assert v2['reset'] == 0x101004
    assert 'merged' not in v2
    assert v2['fields'][0]['reset'] == 0x10
    assert 'dft_value' not in v2['fields'][0]
    assert v2['fields'][1]['reset'] == 0x1004
    assert 'merged' not in status
    assert 'dft_value' not in status['fields'][0]
    assert {'W1_PULSE', 'WC_INT'} <= {field['access'] for field in event['fields']}
    assert not any('写入语义待确认' in warning for warning in preview['warnings'])
    assert parse_register_file('v2-only.xlsx', content, ['V2_EXAMPLE'])['count'] == 2
    for selected in [[], ['missing'], ['说明']]:
        with pytest.raises(RegisterFileError):
            parse_register_file('mixed.xlsx', content, selected)
    # 导出再导入保留两类初始值，但不保留不需要的两列。
    assert parse_register_file('roundtrip.json', json.dumps({'registers': preview['registers']}).encode())['registers'] == preview['registers']


def test_unused_columns_are_not_parsed_or_persisted(tmp_path):
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(versioned_workbook()))
    sheet = workbook['V2_EXAMPLE']
    sheet['C5'] = '不需要解析的任意标记'
    sheet['J5'] = '=1+1'
    content = io.BytesIO()
    workbook.save(content)
    workbook.close()
    registers = parse_register_file('v2.xlsx', content.getvalue(), ['V2_EXAMPLE'])['registers']
    # 本轮早期 JSON 也允许读入，但立即丢弃两项元数据。
    registers[0]['merged'] = True
    registers[0]['fields'][0]['dft_value'] = 'discard me'
    repository = Repository(tmp_path / 'metadata.db')
    repository.save_register_catalog(1, registers)
    service = StudioService(repository.path)
    try:
        exported = service.register_catalog()['registers']
        assert exported[0]['reset'] == 0x101004
        assert 'merged' not in exported[0]
        assert 'dft_value' not in exported[0]['fields'][0]
        assert 'merged' not in service.repository.register_catalog()['registers'][0]
        assert 'dft_value' not in service.repository.register_catalog()['registers'][0]['fields'][0]
    finally:
        service.close()


def test_self_clear_aliases_preserve_type_and_allow_explicit_write(tmp_path):
    service = StudioService(tmp_path / 'special.db')
    try:
        registers = parse_register_file('v1.xlsx', versioned_workbook(), ['V1_EXAMPLE'])['registers']
        service.import_registers(RegisterImport(expected_revision=service.catalog_revision, registers=registers))
        connection_id = service.snapshot()['connections'][0]['id']
        service.connect(connection_id)
        for bit in (0, 1):
            result = service.write_register(connection_id, RegisterWrite(address='0x5004', value=1 << bit, mask=1 << bit))
            assert result['verified'] is None
            assert result['unverified_mask'] == 1 << bit
        assert service.read_register(connection_id, 0x5004)['value'] == 0
    finally:
        service.close()
